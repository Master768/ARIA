import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) for a table cell in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Apply background color to a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    """
    Set cell borders: top, bottom, left, right.
    Format: top={"sz": 4, "val": "single", "color": "CCCCCC"}
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_footer_page_number(run):
    """Appends dynamic Word XML PAGE field into a text run."""
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def create_base_document():
    doc = docx.Document()
    
    # Page setup: A4, 1-inch margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.footer_distance = Inches(0.5)
    
    # Suppress page number on Cover Page, show on subsequent pages
    section.different_first_page_header_footer = True
    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run()
    r_foot.font.name = 'Times New Roman'
    r_foot.font.size = Pt(11)
    r_foot.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    add_footer_page_number(r_foot)
    
    # Configure Normal Style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    return doc

def add_title(doc, text, space_before=12, space_after=18):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Deep Slate/Navy
    return p

def add_heading_1(doc, text, space_before=18, space_after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Academic Royal Navy
    return p

def add_heading_2(doc, text, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_heading_3(doc, text, space_before=10, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    return p

def add_bullet_point(doc, text, bold_title=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(4)
    if bold_title:
        r_title = p.add_run(bold_title + " ")
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(12)
        r_title.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def add_callout_box(doc, speaker, role, text):
    """Formatted callout container for Council debates and architectural invariants."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_shading(cell, "F8FAFC") # Soft Slate / Ice
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Left accent border
    set_cell_border(cell, 
                    left={"sz": 24, "val": "single", "color": "1E3A8A"},
                    top={"sz": 4, "val": "single", "color": "E2E8F0"},
                    bottom={"sz": 4, "val": "single", "color": "E2E8F0"},
                    right={"sz": 4, "val": "single", "color": "E2E8F0"})
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(2)
    
    r_speaker = p.add_run(f"[{speaker} — {role}]: ")
    r_speaker.font.name = 'Times New Roman'
    r_speaker.font.size = Pt(11)
    r_speaker.font.bold = True
    r_speaker.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_text = p.add_run(f'"{text}"')
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(11)
    r_text.font.italic = True
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, data, col_widths=None):
    """Create publication-grade academic tables with professional styling."""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_row = table.rows[0]
    for i, header_text in enumerate(headers):
        cell = hdr_row.cells[i]
        if col_widths and i < len(col_widths):
            cell.width = Inches(col_widths[i])
        set_cell_shading(cell, "1E3A8A") # Navy
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        set_cell_border(cell, 
                        top={"sz": 6, "val": "single", "color": "1E3A8A"},
                        bottom={"sz": 12, "val": "single", "color": "0F172A"},
                        left={"sz": 4, "val": "single", "color": "CBD5E1"},
                        right={"sz": 4, "val": "single", "color": "CBD5E1"})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header_text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    # Data Rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row.cells[c_idx]
            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])
            set_cell_shading(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            set_cell_border(cell, 
                            top={"sz": 4, "val": "single", "color": "E2E8F0"},
                            bottom={"sz": 4, "val": "single", "color": "E2E8F0"},
                            left={"sz": 4, "val": "single", "color": "E2E8F0"},
                            right={"sz": 4, "val": "single", "color": "E2E8F0"})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            # Left-align first column if text, center-align numbers
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(cell_value))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table
