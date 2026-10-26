import pptx

prs = pptx.Presentation(r'C:\Users\Ravin\OneDrive\Desktop\Orchestra\Dissertation_I_Review_PPT_Template.pptx')

slide = prs.slides[1] # Slide 2
shape = slide.shapes[5]
tf = shape.text_frame

print("Left:", shape.left, "Top:", shape.top, "Width:", shape.width, "Height:", shape.height)
print("Word wrap:", tf.word_wrap)

for i, p in enumerate(tf.paragraphs):
    print(f"P{i}: font={p.font.name}, size={p.font.size.pt if p.font.size else None}, bold={p.font.bold}, level={p.level}, text='{p.text}'")
    for r in p.runs:
        print(f"   Run: text='{r.text}', font={r.font.name}, size={r.font.size.pt if r.font.size else None}")
