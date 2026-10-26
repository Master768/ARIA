import pptx

prs = pptx.Presentation(r'C:\Users\Ravin\OneDrive\Desktop\Orchestra\Dissertation_I_Review_PPT_Template.pptx')

for i, slide in enumerate(prs.slides):
    title = 'NO_TITLE'
    content_shape_idx = None
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if s_idx == 0 and i > 0:
                title = txt
            if shape.text_frame.text.startswith('•') or 'Write a concise abstract' in shape.text_frame.text or 'Research context' in shape.text_frame.text or s_idx == 5:
                content_shape_idx = s_idx
    print(f"Slide {i+1}: Title='{title}', ContentShapeIdx={content_shape_idx}")
