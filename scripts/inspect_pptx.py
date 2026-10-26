import pptx

prs = pptx.Presentation(r'C:\Users\Ravin\OneDrive\Desktop\Orchestra\Dissertation_I_Review_PPT_Template.pptx')

print(f'Total slides: {len(prs.slides)}')

for i, slide in enumerate(prs.slides):
    title = 'NO_TITLE'
    if slide.shapes.title and slide.shapes.title.text.strip():
        title = slide.shapes.title.text.strip()
    print(f'\n--- SLIDE {i+1} [Title: {title}] ---')
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            tf = shape.text_frame
            for p_idx, p in enumerate(tf.paragraphs):
                txt = p.text.strip()
                if txt:
                    print(f'   Shape {s_idx} P{p_idx}: {txt}')
        elif shape.has_table:
            t = shape.table
            print(f'   Shape {s_idx} [TABLE {len(t.rows)}x{len(t.columns)}]')
            for r_idx, row in enumerate(t.rows):
                row_vals = [cell.text.strip() for cell in row.cells]
                print(f'     Row {r_idx}: {" | ".join(row_vals)}')
