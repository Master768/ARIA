import os
import json
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

class ReportWriter:
    def __init__(self):
        self.reports_dir = os.path.abspath("reports")
        os.makedirs(self.reports_dir, exist_ok=True)
        self.manifest_path = os.path.join(self.reports_dir, "manifest.json")
        self._init_manifest()

    def _init_manifest(self):
        """Initializes manifest.json if it doesn't exist."""
        if not os.path.exists(self.manifest_path):
            with open(self.manifest_path, "w", encoding="utf-8") as f:
                json.dump({"reports": []}, f, indent=2)

    def _load_manifest(self) -> dict:
        with open(self.manifest_path, "r", encoding="utf-8") as f:
            return json.load(f)

    def _save_manifest(self, data: dict):
        with open(self.manifest_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)

    def compile_report(self, task_state: dict) -> str:
        """
        Compiles the validated findings from the task state into a structured Markdown draft report.
        Also registers the report draft in manifest.json.
        """
        # Determine file name parameters
        timestamp_str = datetime.now().strftime("%Y-%m-%d_%H-%M")
        original_q = task_state.get("original_question", "data_analysis")
        
        # Clean question string to make a safe topic slug
        topic_slug = "".join(c if c.isalnum() else "_" for c in original_q.lower())
        topic_slug = "_".join(filter(None, topic_slug.split("_")))[:30]
        
        file_name = f"{timestamp_str}_{topic_slug}_DRAFT.md"
        file_path = os.path.join(self.reports_dir, file_name)
        
        # Build markdown content
        md_lines = []
        md_lines.append(f"# Analysis Report: {original_q}")
        md_lines.append(f"\n**Date:** {task_state.get('timestamp', timestamp_str)}")
        md_lines.append(f"**Status:** {task_state.get('status', 'DRAFT')}\n")
        md_lines.append("## Executive Summary")
        md_lines.append(task_state.get("final_summary", "No summary provided."))
        md_lines.append("\n## Detailed Findings")
        
        for task in task_state.get("sub_tasks", []):
            md_lines.append(f"### Sub-task {task.get('id')}: {task.get('description')}")
            md_lines.append(f"- **Query:** {task.get('query')}")
            
            agent_res = task.get("data_agent_result", {})
            if isinstance(agent_res, dict):
                output_content = agent_res.get("output", "No output generated.")
            else:
                output_content = str(agent_res)
            
            md_lines.append(f"- **Result:** {output_content}")
            md_lines.append(f"- **Critic Verification:** {task.get('critic_validation', 'PENDING')}")
            md_lines.append("")
            
        md_content = "\n".join(md_lines)
        
        # Save markdown report
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(md_content)

        # Save task state copy
        state_file_name = f"{timestamp_str}_{topic_slug}_state.json"
        state_file_path = os.path.join(self.reports_dir, state_file_name)
        with open(state_file_path, "w", encoding="utf-8") as f:
            json.dump(task_state, f, indent=2)
            
        # Register draft in manifest.json
        manifest = self._load_manifest()
        # Prevent duplicate entries for the same path
        manifest["reports"] = [r for r in manifest["reports"] if r.get("md_path") != file_name]
        
        manifest["reports"].append({
            "timestamp": task_state.get("timestamp", datetime.now().isoformat()),
            "original_question": original_q,
            "status": "DRAFT",
            "validation_status": task_state.get("status"), # COMPLETE / NEEDS_REVIEW
            "md_path": file_name,
            "state_path": state_file_name,
            "docx_path": None
        })
        self._save_manifest(manifest)
        
        return md_content

    def convert_to_docx(self, markdown_content: str, docx_path: str):
        """
        Parses Markdown content and writes a professionally styled Word document
        with custom page margins, typography, headers, footers, and a metadata banner.
        """
        doc = docx.Document()
        
        # 1. Page Setup: 1 inch margins all around
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
            
        # 2. Configure Typography Styles
        # Normal (Body Text)
        style_normal = doc.styles['Normal']
        font_normal = style_normal.font
        font_normal.name = 'Calibri'
        font_normal.size = Pt(11)
        font_normal.color.rgb = RGBColor(0x33, 0x33, 0x33)
        style_normal.paragraph_format.line_spacing = 1.15
        style_normal.paragraph_format.space_after = Pt(6)
        
        # Headings
        style_h1 = doc.styles['Heading 1']
        font_h1 = style_h1.font
        font_h1.name = 'Calibri'
        font_h1.size = Pt(20)
        font_h1.bold = True
        font_h1.color.rgb = RGBColor(0x1F, 0x4E, 0x79) # Navy Blue
        style_h1.paragraph_format.space_before = Pt(12)
        style_h1.paragraph_format.space_after = Pt(6)
        
        style_h2 = doc.styles['Heading 2']
        font_h2 = style_h2.font
        font_h2.name = 'Calibri'
        font_h2.size = Pt(14)
        font_h2.bold = True
        font_h2.color.rgb = RGBColor(0x2E, 0x75, 0xB6) # Steel Blue
        style_h2.paragraph_format.space_before = Pt(12)
        style_h2.paragraph_format.space_after = Pt(4)
        
        style_h3 = doc.styles['Heading 3']
        font_h3 = style_h3.font
        font_h3.name = 'Calibri'
        font_h3.size = Pt(12)
        font_h3.bold = True
        font_h3.color.rgb = RGBColor(0x59, 0x59, 0x59) # Dark Gray
        style_h3.paragraph_format.space_before = Pt(6)
        style_h3.paragraph_format.space_after = Pt(2)

        # 3. Parse Metadata for Banner Table
        title = "Data Analysis Report"
        date_str = datetime.now().strftime("%Y-%m-%d")
        status_str = "APPROVED"
        
        lines = markdown_content.splitlines()
        for line in lines:
            line_strip = line.strip()
            if line_strip.startswith("# "):
                title = line_strip[2:]
            elif line_strip.startswith("**Date:**"):
                date_str = line_strip[9:].strip().replace("**", "")
            elif line_strip.startswith("**Status:**"):
                status_str = line_strip[11:].strip().replace("**", "")

        # 4. Create Metadata Banner Table (1 row, 2 columns)
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(4.5)
        table.columns[1].width = Inches(2.0)
        
        # Shade and Border cells
        def set_cell_shading(cell, color_hex):
            shd_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
            cell._tc.get_or_add_tcPr().append(shd_elm)
            
        def set_cell_left_border(cell, color_hex, size="36"):
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/></w:tcBorders>')
            cell._tc.get_or_add_tcPr().append(tcBorders)

        cell_left = table.cell(0, 0)
        cell_right = table.cell(0, 1)
        
        set_cell_shading(cell_left, "F2F5F8")
        set_cell_shading(cell_right, "F2F5F8")
        set_cell_left_border(cell_left, "1F4E79", size="36")
        
        p_title = cell_left.paragraphs[0]
        run_title = p_title.add_run(title)
        run_title.font.size = Pt(16)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
        p_meta = cell_right.paragraphs[0]
        p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_meta.add_run("DATE:\n").bold = True
        p_meta.add_run(f"{date_str}\n\n")
        p_meta.add_run("STATUS:\n").bold = True
        p_meta.add_run(status_str).font.color.rgb = RGBColor(0x10, 0x7C, 0x41) if status_str == "APPROVED" else RGBColor(0xA6, 0x1C, 0x1C)
        
        # Spacing element
        doc.add_paragraph().paragraph_format.space_before = Pt(12)

        # 5. Populate Document Content
        def add_formatted_paragraph(doc, text, style=None):
            p = doc.add_paragraph(style=style)
            parts = text.split("**")
            for idx, part in enumerate(parts):
                is_bold = (idx % 2 == 1)
                p.add_run(part).bold = is_bold
            return p

        for line in lines:
            line_strip = line.strip()
            if not line_strip:
                continue
                
            # Skip metadata items that were placed in the banner
            if line_strip.startswith("# ") or line_strip.startswith("**Date:**") or line_strip.startswith("**Status:**"):
                continue
                
            if line_strip.startswith("## "):
                doc.add_heading(line_strip[3:], level=2)
            elif line_strip.startswith("### "):
                doc.add_heading(line_strip[4:], level=3)
            elif line_strip.startswith("- ") or line_strip.startswith("* "):
                add_formatted_paragraph(doc, line_strip[2:], style="List Bullet")
            else:
                add_formatted_paragraph(doc, line_strip)
                
        # 6. Add Footer
        for section in doc.sections:
            footer = section.footer
            f_p = footer.paragraphs[0]
            f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            f_run = f_p.add_run("Orchestra Analytics Report Studio  |  Confidential  |  Page ")
            f_run.font.size = Pt(9)
            f_run.font.italic = True
            f_run.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
            
            # XML page number field
            fldSimple = OxmlElement('w:fldSimple')
            fldSimple.set(qn('w:instr'), 'PAGE')
            f_p._p.append(fldSimple)

        doc.save(docx_path)

    def update_manifest_status(self, draft_filename: str, new_status: str, new_filename: str = None, docx_filename: str = None):
        """Updates report status and file paths inside reports/manifest.json."""
        manifest = self._load_manifest()
        found = False
        for report in manifest["reports"]:
            if report.get("md_path") == draft_filename:
                report["status"] = new_status
                if new_filename:
                    report["md_path"] = new_filename
                if docx_filename:
                    report["docx_path"] = docx_filename
                found = True
                break
        if found:
            self._save_manifest(manifest)
