import os
import json
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))

class ReportWriter:
    """Tier 5 Executive Synthesis Intelligence: ReportWriter."""
    def __init__(self, reports_dir: str = None):
        if reports_dir:
            self.reports_dir = os.path.abspath(reports_dir)
        else:
            self.reports_dir = os.path.join(PROJECT_ROOT, "reports")
        os.makedirs(self.reports_dir, exist_ok=True)
        self.manifest_path = os.path.join(self.reports_dir, "manifest.json")
        self._init_manifest()

    def _init_manifest(self):
        if not os.path.exists(self.manifest_path):
            with open(self.manifest_path, "w", encoding="utf-8") as f:
                json.dump({"reports": []}, f, indent=2)

    def _load_manifest(self) -> dict:
        with open(self.manifest_path, "r", encoding="utf-8") as f:
            return json.load(f)

    def _save_manifest(self, data: dict):
        with open(self.manifest_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)

    def compile_report(self, task_state: dict) -> str:
        timestamp_str = datetime.now().strftime("%Y-%m-%d_%H-%M")
        original_q = task_state.get("original_question", "data_analysis")
        
        topic_slug = "".join(c if c.isalnum() else "_" for c in original_q.lower())
        topic_slug = "_".join(filter(None, topic_slug.split("_")))[:30]
        
        file_name = f"{timestamp_str}_{topic_slug}_DRAFT.md"
        file_path = os.path.join(self.reports_dir, file_name)
        
        md_lines = []
        md_lines.append(f"# Analysis Report: {original_q}")
        md_lines.append(f"\n**Date:** {task_state.get('timestamp', timestamp_str)}")
        md_lines.append(f"**Status:** {task_state.get('status', 'DRAFT')}\n")
        md_lines.append("## Executive Summary")
        md_lines.append(task_state.get("final_summary", "No summary provided."))
        md_lines.append("\n## Detailed Findings")
        
        for task in task_state.get("sub_tasks", []):
            md_lines.append(f"### Sub-task {task.get('id')}: {task.get('description')}")
            md_lines.append(f"- **Query:** {task.get('query')}")
            
            agent_res = task.get("data_agent_result", {})
            output_content = agent_res.get("output", "No output generated.") if isinstance(agent_res, dict) else str(agent_res)
            
            md_lines.append(f"- **Result:** {output_content}")
            md_lines.append(f"- **Critic Verification:** {task.get('critic_validation', 'PENDING')}")
            md_lines.append("")
            
        md_content = "\n".join(md_lines)
        
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(md_content)

        state_file_name = f"{timestamp_str}_{topic_slug}_state.json"
        state_file_path = os.path.join(self.reports_dir, state_file_name)
        with open(state_file_path, "w", encoding="utf-8") as f:
            json.dump(task_state, f, indent=2)
            
        manifest = self._load_manifest()
        manifest["reports"] = [r for r in manifest["reports"] if r.get("md_path") != file_name]
        manifest["reports"].append({
            "timestamp": task_state.get("timestamp", datetime.now().isoformat()),
            "original_question": original_q,
            "status": "DRAFT",
            "validation_status": task_state.get("status"),
            "md_path": file_name,
            "state_path": state_file_name,
            "docx_path": None
        })
        self._save_manifest(manifest)
        return md_content

    def convert_to_docx(self, markdown_content: str, docx_path: str):
        doc = docx.Document()
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        style_normal = doc.styles['Normal']
        font_normal = style_normal.font
        font_normal.name = 'Calibri'
        font_normal.size = Pt(11)

        lines = markdown_content.splitlines()
        for line in lines:
            line_strip = line.strip()
            if not line_strip:
                continue
            if line_strip.startswith("## "):
                doc.add_heading(line_strip[3:], level=2)
            elif line_strip.startswith("### "):
                doc.add_heading(line_strip[4:], level=3)
            else:
                doc.add_paragraph(line_strip)

        doc.save(docx_path)
